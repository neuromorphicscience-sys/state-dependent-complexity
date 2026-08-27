#!/usr/bin/env python3
"""Build the Supplementary Table 1 writing and submission package."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle


HERE = Path(__file__).resolve().parent
PLOT_ROOT = HERE.parents[2]
SOURCE = (
    PLOT_ROOT
    / "Manuscript_Figures_Nature_DataFirst_V5_3"
    / "03_SUPPLEMENTARY"
    / "S01"
    / "S01_c__Zero-variance_structural-variable_audit"
    / "source_data"
    / "01_S5A_18_PFC_zero_variance_audit__pfc_feature_qc.csv"
)

TITLE = (
    "Supplementary Table 1 | Quality-control audit of structural variables "
    "in the PFC atlas cohort"
)
CAPTION = (
    "Supplementary Table 1. Quality-control audit of structural variables in "
    "the PFC atlas cohort. For each archived field, the table reports the "
    "number of observations, number of unique values and descriptive range. "
    "Identifier and categorical-code fields are retained for provenance but "
    "are not interpreted as continuous biological quantities. Dendrite length "
    "and dendrite nodes have zero variance in the archived cohort and are "
    "therefore uninformative for variance-based analyses. No inferential "
    "statistics are presented."
)
NOTES = [
    "n is the number of archived observations; unique is the number of distinct archived values.",
    "Zero variance is defined here by unique = 1 and s.d. = 0.",
    "Numeric summaries use the archived source scale; no units or transformations were reassigned.",
    "Complete unmodified values are retained in the Raw source worksheet and the accompanying CSV.",
]

TEAL = "2A9D8F"
TEAL_DARK = "227D73"
PALE_TEAL = "E7F4F2"
PALE_RED = "FCE8E6"
PALE_GRAY = "F2F4F5"
INK = "263238"
MID = "657178"
WHITE = "FFFFFF"


DISPLAY_NAME = {
    "archive_id": "Archive ID",
    "axon_length": "Axon length",
    "axon_nodes": "Axon nodes",
    "bbox_dx": "Bounding-box extent, x",
    "bbox_dy": "Bounding-box extent, y",
    "bbox_dz": "Bounding-box extent, z",
    "bbox_volume": "Bounding-box volume",
    "branch_points": "Branch points",
    "cell_id": "Cell ID",
    "dendrite_length": "Dendrite length",
    "dendrite_nodes": "Dendrite nodes",
    "endpoints": "Endpoints",
    "max_euclidean_radius": "Maximum Euclidean radius",
    "max_path_length": "Maximum path length",
    "n_nodes": "Total nodes",
    "node_type_entropy": "Node-type entropy",
    "path_tortuosity_proxy": "Path-tortuosity proxy",
    "projection_subtype": "Projection subtype",
    "radius_gyration": "Radius of gyration",
    "soma_nodes": "Soma nodes",
    "soma_x": "Soma coordinate, x",
    "soma_y": "Soma coordinate, y",
    "soma_z": "Soma coordinate, z",
    "spatial_entropy_4": "Spatial entropy (4)",
    "spatial_entropy_6": "Spatial entropy (6)",
    "total_length": "Total length",
}

IDENTIFIERS = {"archive_id", "cell_id"}
CATEGORICAL = {"projection_subtype"}
COUNTS = {
    "axon_nodes", "branch_points", "dendrite_nodes", "endpoints",
    "n_nodes", "soma_nodes",
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def variable_type(name: str) -> str:
    if name in IDENTIFIERS:
        return "Identifier"
    if name in CATEGORICAL:
        return "Categorical code"
    if name in COUNTS:
        return "Count"
    return "Continuous"


def format_value(value: float | int) -> str:
    value = float(value)
    if value == 0:
        return "0"
    if abs(value) >= 100000 or abs(value) < 0.001:
        return f"{value:.3e}"
    if abs(value) >= 1000:
        return f"{value:,.1f}"
    if abs(value) >= 10:
        return f"{value:.2f}"
    return f"{value:.3f}"


def prepare_table(raw: pd.DataFrame) -> pd.DataFrame:
    records = []
    for row in raw.itertuples(index=False):
        kind = variable_type(row.feature)
        zero = bool(row.zero_variance)
        interpret_numeric = kind not in {"Identifier", "Categorical code"}
        if zero:
            interpretation = "Zero variance; not informative for variance-based analyses"
        elif kind == "Identifier":
            interpretation = "Identifier; descriptive numeric moments not interpreted"
        elif kind == "Categorical code":
            interpretation = "Categorical code; descriptive numeric moments not interpreted"
        else:
            interpretation = "Non-zero variance"
        records.append({
            "Variable": DISPLAY_NAME.get(row.feature, row.feature.replace("_", " ").title()),
            "Type": kind,
            "n": int(row.n),
            "Unique": int(row.n_unique),
            "Mean": format_value(row.mean) if interpret_numeric else "NA",
            "s.d.": format_value(row.std) if interpret_numeric else "NA",
            "Minimum": format_value(row.min) if interpret_numeric else "NA",
            "Maximum": format_value(row.max) if interpret_numeric else "NA",
            "Zero variance": "Yes" if zero else "No",
            "QC interpretation": interpretation,
            "source_field": row.feature,
        })
    return pd.DataFrame.from_records(records)


def build_xlsx(table: pd.DataFrame, raw: pd.DataFrame, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Table 1"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:J1")
    ws["A1"] = TITLE
    ws["A1"].font = Font(name="Arial", size=12, bold=True, color=INK)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:J2")
    ws["A2"] = CAPTION
    ws["A2"].font = Font(name="Arial", size=9, color=MID)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 62

    headers = list(table.columns[:-1])
    header_row = 4
    for col, header in enumerate(headers, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL_DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin = Side(style="thin", color="D9DEE1")
    for r_idx, record in enumerate(table[headers].itertuples(index=False), header_row + 1):
        zero = record[8] == "Yes"
        neutral = record[1] in {"Identifier", "Categorical code"}
        fill = PALE_RED if zero else (PALE_GRAY if neutral else (PALE_TEAL if r_idx % 2 else WHITE))
        for c_idx, value in enumerate(record, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.font = Font(name="Arial", size=8.5, color=INK,
                             bold=(zero and c_idx in {1, 9}))
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(
                horizontal="left" if c_idx in {1, 2, 10} else "center",
                vertical="center", wrap_text=(c_idx in {1, 2, 10}),
            )
        ws.row_dimensions[r_idx].height = 31 if zero or neutral else 23

    note_start = header_row + len(table) + 2
    for i, note in enumerate(NOTES, 1):
        ws.merge_cells(start_row=note_start + i - 1, start_column=1,
                       end_row=note_start + i - 1, end_column=10)
        cell = ws.cell(note_start + i - 1, 1, f"Note {i}. {note}")
        cell.font = Font(name="Arial", size=8, color=MID)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = [29, 17, 10, 11, 14, 14, 14, 14, 12, 51]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:J{header_row + len(table)}"
    ws.print_title_rows = "1:4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = .3
    ws.page_margins.right = .3
    ws.page_margins.top = .45
    ws.page_margins.bottom = .45
    ws.oddFooter.center.text = "Supplementary Table 1"
    ws.oddFooter.right.text = "Page &P of &N"

    raw_ws = wb.create_sheet("Raw source")
    raw_ws.sheet_view.showGridLines = False
    for c_idx, header in enumerate(raw.columns, 1):
        c = raw_ws.cell(1, c_idx, header)
        c.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=TEAL_DARK)
        c.alignment = Alignment(horizontal="center")
    for r_idx, row in enumerate(raw.itertuples(index=False), 2):
        for c_idx, value in enumerate(row, 1):
            c = raw_ws.cell(r_idx, c_idx, value)
            c.font = Font(name="Arial", size=8.5, color=INK)
            c.alignment = Alignment(horizontal="left" if c_idx == 1 else "right")
            if c_idx in {4, 5, 6, 7}:
                c.number_format = "0.000E+00"
            if c_idx == 8:
                c.fill = PatternFill("solid", fgColor=PALE_RED if bool(value) else WHITE)
    raw_ws.freeze_panes = "A2"
    raw_ws.auto_filter.ref = f"A1:H{len(raw) + 1}"
    raw_widths = [27, 11, 13, 17, 17, 17, 17, 15]
    for idx, width in enumerate(raw_widths, 1):
        raw_ws.column_dimensions[get_column_letter(idx)].width = width

    readme = wb.create_sheet("Readme")
    readme.sheet_view.showGridLines = False
    entries = [
        ("Deliverable", "Supplementary Table 1"),
        ("Status", "Draft ready for manuscript writing"),
        ("Source", str(SOURCE.relative_to(PLOT_ROOT))),
        ("Rows", str(len(raw))),
        ("Archived observations", f"{int(raw['n'].max()):,}"),
        ("Zero-variance fields", ", ".join(raw.loc[raw.zero_variance, 'feature'])),
        ("Interpretation boundary", "This is a descriptive QC table; it contains no inferential statistics."),
    ]
    for r_idx, (key, value) in enumerate(entries, 1):
        readme.cell(r_idx, 1, key).font = Font(name="Arial", size=9, bold=True, color=INK)
        readme.cell(r_idx, 2, value).font = Font(name="Arial", size=9, color=INK)
        readme.cell(r_idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 27
    readme.column_dimensions["B"].width = 105

    wb.properties.title = TITLE
    wb.properties.subject = "PFC structural-variable quality-control audit"
    wb.properties.creator = "Neural Science manuscript figure workflow"
    wb.save(path)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def build_docx(table: pd.DataFrame, path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(11)
    section.bottom_margin = Mm(11)
    section.left_margin = Mm(11)
    section.right_margin = Mm(11)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(8)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(TITLE)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(INK)

    intro = doc.add_paragraph(CAPTION)
    intro.paragraph_format.space_after = Pt(6)
    intro.paragraph_format.line_spacing = 1.05
    for run in intro.runs:
        run.font.name = "Arial"
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string(MID)

    headers = list(table.columns[:-1])
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.autofit = False
    set_repeat_table_header(tbl.rows[0])
    widths_mm = [37, 20, 14, 17, 21, 21, 21, 21, 18, 83]
    for idx, (cell, header, width) in enumerate(zip(tbl.rows[0].cells, headers, widths_mm)):
        cell.width = Mm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, TEAL_DARK)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(6.7)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for record in table[headers].itertuples(index=False):
        row = tbl.add_row()
        zero = record[8] == "Yes"
        neutral = record[1] in {"Identifier", "Categorical code"}
        fill = PALE_RED if zero else (PALE_GRAY if neutral else WHITE)
        for idx, (cell, value, width) in enumerate(zip(row.cells, record, widths_mm)):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, fill)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx in {0, 1, 9} else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(value))
            r.font.name = "Arial"
            r.font.size = Pt(6.3)
            r.font.color.rgb = RGBColor.from_string(INK)
            r.bold = bool(zero and idx in {0, 8})

    notes_head = doc.add_paragraph()
    notes_head.paragraph_format.space_before = Pt(5)
    notes_head.paragraph_format.space_after = Pt(1)
    r = notes_head.add_run("Notes")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(7.2)
    for idx, note in enumerate(NOTES, 1):
        p = doc.add_paragraph(f"{idx}. {note}")
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Mm(3)
        p.paragraph_format.first_line_indent = Mm(-3)
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(6.7)
            r.font.color.rgb = RGBColor.from_string(MID)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Supplementary Table 1")
    r.font.name = "Arial"
    r.font.size = Pt(7)
    r.font.color.rgb = RGBColor.from_string(MID)

    props = doc.core_properties
    props.title = TITLE
    props.subject = "PFC structural-variable quality-control audit"
    props.author = "Neural Science manuscript workflow"
    doc.save(path)


def build_pdf(table: pd.DataFrame, path: Path) -> None:
    arial_value = os.environ.get("ARIAL_FONT")
    arial_bold_value = os.environ.get("ARIAL_BOLD_FONT")
    arial = Path(arial_value).expanduser() if arial_value else None
    arial_bold = Path(arial_bold_value).expanduser() if arial_bold_value else None
    if arial and arial_bold and arial.is_file() and arial_bold.is_file():
        pdfmetrics.registerFont(TTFont("Arial", str(arial)))
        pdfmetrics.registerFont(TTFont("Arial-Bold", str(arial_bold)))
        body_font, bold_font = "Arial", "Arial-Bold"
    else:
        body_font, bold_font = "Helvetica", "Helvetica-Bold"

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TableTitle", parent=styles["Title"], fontName=bold_font,
        fontSize=10.5, leading=12.5, textColor=colors.HexColor(f"#{INK}"),
        alignment=TA_LEFT, spaceAfter=4,
    )
    caption_style = ParagraphStyle(
        "Caption", parent=styles["BodyText"], fontName=body_font,
        fontSize=7.2, leading=8.6, textColor=colors.HexColor(f"#{MID}"),
        alignment=TA_LEFT, spaceAfter=6,
    )
    cell_style = ParagraphStyle(
        "Cell", parent=styles["BodyText"], fontName=body_font,
        fontSize=6.0, leading=7.0, textColor=colors.HexColor(f"#{INK}"),
        alignment=TA_LEFT,
    )
    center_style = ParagraphStyle(
        "CellCenter", parent=cell_style, alignment=TA_CENTER,
    )
    header_style = ParagraphStyle(
        "Header", parent=center_style, fontName=bold_font,
        fontSize=6.1, leading=7.1, textColor=colors.white,
    )
    note_style = ParagraphStyle(
        "Note", parent=styles["BodyText"], fontName=body_font,
        fontSize=6.5, leading=7.8, textColor=colors.HexColor(f"#{MID}"),
        alignment=TA_LEFT, spaceAfter=1.5,
    )

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(body_font, 6.5)
        canvas.setFillColor(colors.HexColor(f"#{MID}"))
        canvas.drawString(12 * mm, 6.5 * mm, "Supplementary Table 1")
        canvas.drawRightString(285 * mm, 6.5 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(path), pagesize=landscape(A4),
        leftMargin=11 * mm, rightMargin=11 * mm,
        topMargin=10 * mm, bottomMargin=11 * mm,
        title=TITLE, author="Neural Science manuscript workflow",
    )
    story = [Paragraph(TITLE, title_style), Paragraph(CAPTION, caption_style)]
    headers = list(table.columns[:-1])
    rows = [[Paragraph(h, header_style) for h in headers]]
    zero_rows = []
    neutral_rows = []
    for idx, record in enumerate(table[headers].itertuples(index=False), 1):
        if record[8] == "Yes":
            zero_rows.append(idx)
        elif record[1] in {"Identifier", "Categorical code"}:
            neutral_rows.append(idx)
        rows.append([
            Paragraph(str(value), cell_style if col in {0, 1, 9} else center_style)
            for col, value in enumerate(record)
        ])
    widths = [37, 20, 14, 17, 21, 21, 21, 21, 18, 83]
    table_obj = LongTable(rows, colWidths=[w * mm for w in widths], repeatRows=1,
                          hAlign="LEFT")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{TEAL_DARK}")),
        ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#D9DEE1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.4),
        ("TOPPADDING", (0, 0), (-1, -1), 2.6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.6),
    ]
    for idx in zero_rows:
        commands.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor(f"#{PALE_RED}")))
    for idx in neutral_rows:
        commands.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor(f"#{PALE_GRAY}")))
    table_obj.setStyle(TableStyle(commands))
    story.append(table_obj)
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>Notes</b>", note_style))
    story.extend(Paragraph(f"{idx}. {note}", note_style) for idx, note in enumerate(NOTES, 1))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def main() -> None:
    HERE.mkdir(parents=True, exist_ok=True)
    raw = pd.read_csv(SOURCE)
    expected = {"feature", "n", "n_unique", "mean", "std", "min", "max", "zero_variance"}
    if set(raw.columns) != expected:
        raise RuntimeError(f"Unexpected source schema: {list(raw.columns)}")
    table = prepare_table(raw)

    source_copy = HERE / "Supplementary_Table_1_source.csv"
    shutil.copy2(SOURCE, source_copy)
    caption_path = HERE / "Supplementary_Table_1_caption.txt"
    caption_path.write_text(CAPTION + "\n\nNotes:\n" + "\n".join(
        f"{idx}. {note}" for idx, note in enumerate(NOTES, 1)
    ) + "\n", encoding="utf-8")

    xlsx = HERE / "Supplementary_Table_1.xlsx"
    docx = HERE / "Supplementary_Table_1.docx"
    pdf = HERE / "Supplementary_Table_1.pdf"
    build_xlsx(table, raw, xlsx)
    build_docx(table, docx)
    build_pdf(table, pdf)

    artifacts = [
        xlsx, docx, pdf, source_copy, caption_path,
        HERE / "README.md", HERE / "QA_REPORT.md", Path(__file__),
    ]
    manifest = {
        "package": "Supplementary_Table_1",
        "status": "DRAFT_READY_FOR_WRITING",
        "title": TITLE,
        "source": str(SOURCE.relative_to(PLOT_ROOT)),
        "source_rows": int(len(raw)),
        "source_columns": list(raw.columns),
        "zero_variance_fields": raw.loc[raw.zero_variance, "feature"].tolist(),
        "artifacts": {
            p.name: {"sha256": sha256(p), "bytes": p.stat().st_size}
            for p in artifacts
        },
    }
    (HERE / "DELIVERY_MANIFEST.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
